import re
import sys
import time
from pathlib import Path

import chromadb
import ollama
from docx import Document
from pypdf import PdfReader
from sentence_transformers import SentenceTransformer

# SETTINGS

EMBED_MODEL = "all-MiniLM-L6-v2"
LLM_MODEL = "llama3.2"
CHUNK_STRATEGY = "fixed"
CHUNK_SIZE = 500
CHUNK_OVERLAP = 100
TOP_K = 3
MAX_DISTANCE = 0.75
SUPPORTED = (".pdf", ".docx", ".txt")
NOT_FOUND = "I could not find this in the uploaded documents."

SYSTEM_PROMPT = (
    "You answer questions using ONLY the context given by the user. "
    f'If the answer is not in the context, reply exactly: "{NOT_FOUND}" '
    "Be concise. Cite the passages you used like [1], [2]."
)

#SETUP OBJECTS

embedder = SentenceTransformer(EMBED_MODEL)
client = chromadb.PersistentClient(path="chroma_db")
collection = client.get_or_create_collection(
    "docs_" + re.sub(r"[^a-zA-Z0-9]+", "_", EMBED_MODEL),
    metadata={"hnsw:space": "cosine"},
)

#STEP 1 - PARSERS

def clean(text: str) -> str:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

def load_txt(path):
    return [{"text": Path(path).read_text(encoding="utf-8", errors="ignore"), "page": 0}]

def load_pdf(path):
    reader = PdfReader(str(path))
    return [{"text": page.extract_text() or "", "page": n} for n, page in enumerate(reader.pages, start=1)]

def load_docx(path):
    doc = Document(str(path))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return [{"text": "\n".join(lines), "page": 0}]

def load_file(path):
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        segments = load_pdf(path)
    elif ext == ".docx":
        segments = load_docx(path)
    elif ext == ".txt":
        segments = load_txt(path)
    else:
        raise ValueError(f"Unsupported file type '{ext}'. Use: {', '.join(SUPPORTED)}")
    segments = [{"text": clean(s["text"]), "page": s["page"]} for s in segments]
    segments = [s for s in segments if s["text"]]
    if not segments:
        raise ValueError(f"No text found in {Path(path).name} (scanned PDF or empty file?)")
    return segments


# STEP 2 - CHUNKING

def chunk_fixed(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    if not 0 <= overlap < size:
        raise ValueError("overlap must be >= 0 and smaller than size")
    chunks, start = [], 0
    while start < len(text):
        piece = text[start:start + size].strip()
        if piece:
            chunks.append(piece)
        if start + size >= len(text):
            break
        start += size - overlap
    return chunks

def chunk_sentences(text: str, size: int = CHUNK_SIZE, overlap: int = 0) -> list[str]:
    sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+|\n{2,}", text) if s.strip()]
    chunks, current = [], ""
    for s in sentences:
        if current and len(current) + len(s) + 1 > size:
            chunks.append(current)
            current = ""
        current = (current + " " + s).strip()
    if current:
        chunks.append(current)
    return chunks

def chunk_text(text, strategy=CHUNK_STRATEGY, size=CHUNK_SIZE, overlap=CHUNK_OVERLAP) -> list[str]:
    if strategy == "fixed":
        return chunk_fixed(text, size, overlap)
    if strategy == "sentence":
        return chunk_sentences(text, size, overlap)
    raise ValueError("strategy must be 'fixed' or 'sentence'")


# STEP 3 - INGEST

def ingest(path, strategy=CHUNK_STRATEGY, size=CHUNK_SIZE, overlap=CHUNK_OVERLAP, col=None) -> int:
    col = col or collection
    name = Path(path).name
    chunks, metadatas = [], []
    for seg in load_file(path):
        for piece in chunk_text(seg["text"], strategy, size, overlap):
            chunks.append(piece)
            metadatas.append({"source": name, "page": seg["page"], "chunk": len(chunks) - 1})

    col.delete(where={"source": name})
    embeddings = embedder.encode(chunks, normalize_embeddings=True).tolist()
    col.upsert(
        ids=[f"{name}-{i}" for i in range(len(chunks))],
        documents=chunks,
        embeddings=embeddings,
        metadatas=metadatas,
    )
    return len(chunks)

def ingest_folder(folder="data", **kwargs) -> dict:
    results = {}
    for p in sorted(Path(folder).iterdir()):
        if p.suffix.lower() in SUPPORTED:
            results[p.name] = ingest(p, **kwargs)
    return results


# STEP 4 - RETRIEVE

def embed_question(question: str):
    if "bge" in EMBED_MODEL.lower():
        question = "Represent this sentence for searching relevant passages: " + question
    return embedder.encode([question], normalize_embeddings=True).tolist()

def retrieve(question: str, k: int = TOP_K, max_distance: float = MAX_DISTANCE, col=None):
    col = col or collection
    total = col.count()
    if total == 0:
        return []
    res = col.query(query_embeddings=embed_question(question), n_results=min(k, total))
    hits = []
    for doc, meta, dist in zip(res["documents"][0], res["metadatas"][0], res["distances"][0]):
        hits.append({"text": doc, "source": meta["source"], "page": meta["page"],
                     "chunk": meta["chunk"], "distance": dist})
    return [h for h in hits if h["distance"] <= max_distance]


# STEP 5 - ASK

def ask(question: str, k: int = TOP_K, max_distance: float = MAX_DISTANCE):
    start = time.time()
    hits = retrieve(question, k, max_distance)

    if not hits:
        return NOT_FOUND, [], time.time() - start

    blocks = []
    for i, h in enumerate(hits, start=1):
        where = f"{h['source']}, page {h['page']}" if h["page"] else h["source"]
        blocks.append(f"[{i}] ({where})\n{h['text']}")
    context = "\n\n".join(blocks)
    prompt = f"Context:\n{context}\n\nQuestion: {question}\nAnswer:"

    try:
        resp = ollama.chat(
            model=LLM_MODEL,
            messages=[{"role": "system", "content": SYSTEM_PROMPT},
                      {"role": "user", "content": prompt}],
            options={"temperature": 0.1},
        )
        answer = resp["message"]["content"].strip()
    except Exception as e:
        answer = f"Could not get an answer from Ollama ({e}). Is Ollama running? Did you run: ollama pull {LLM_MODEL}"
        return answer, hits, time.time() - start

    if NOT_FOUND.lower() in answer.lower():
        hits = []
    return answer, hits, time.time() - start


# COMMAND LINE

if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    cmd, arg = sys.argv[1], " ".join(sys.argv[2:])
    if cmd == "ingest":
        if Path(arg).is_dir():
            for name, n in ingest_folder(arg).items():
                print(f"Ingested {n} chunks from {name}")
        else:
            print(f"Ingested {ingest(arg)} chunks from {Path(arg).name}")
    elif cmd == "ask":
        answer, hits, secs = ask(arg)
        print(f"\nAnswer:\n{answer}\n")
        print("Sources:")
        for i, h in enumerate(hits, 1):
            page = f", page {h['page']}" if h["page"] else ""
            print(f"  [{i}] {h['source']}{page} (chunk {h['chunk']}, distance {h['distance']:.2f})")
        print(f"\nLatency: {secs:.1f}s")
    else:
        sys.exit(__doc__)